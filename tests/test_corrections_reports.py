import io
import uuid
from datetime import date

from openpyxl import load_workbook

from tests.conftest import uid

API = "/api/v1"
ENTRY, EXIT = "2026-03-02T08:00:00Z", "2026-03-02T10:00:00Z"  # exactly 2 hours


def scan(client, kiosk, person):
    return client.post(f"{API}/kiosk/scan", json={"qr_token": person["qr_token"]}, headers=kiosk).json()["result"]


def open_session(client, admin, kiosk, person):
    assert scan(client, kiosk, person) == "ENTRY"
    sessions = client.get(f"{API}/attendance/sessions", params={"person_id": person["id"], "status": "OPEN"},
                          headers=admin).json()
    assert len(sessions) == 1
    return sessions[0]


def correct(client, admin, session_id, reason="test", entry=ENTRY, exit_=EXIT):
    return client.post(f"{API}/attendance/sessions/{session_id}/correct",
                       json={"entry_time": entry, "exit_time": exit_, "reason": reason}, headers=admin)


def test_correction_voids_and_replaces(client, admin, kiosk, make_person, instant_type):
    person = make_person(person_type=instant_type)  # 0 s duplicate window so the person can scan again right away
    old = open_session(client, admin, kiosk, person)
    url = f"{API}/attendance/sessions/{old['id']}/correct"

    # reason is mandatory; times must be ordered; unknown session is 404
    assert correct(client, admin, old["id"], reason="  ").status_code == 422
    assert client.post(url, json={"entry_time": ENTRY, "exit_time": EXIT}, headers=admin).status_code == 422
    assert correct(client, admin, old["id"], entry=EXIT, exit_=ENTRY).status_code == 422
    assert correct(client, admin, uuid.uuid4()).status_code == 404

    r = correct(client, admin, old["id"], reason="Forgot to scan out")
    assert r.status_code == 201, r.text
    new = r.json()
    assert new["status"] == "CLOSED" and new["duration_seconds"] == 7200 and new["work_date"] == "2026-03-02"
    assert [(e["event_type"], e["source"]) for e in new["events"]] == [
        ("ENTRY", "ADMIN_CORRECTION"), ("EXIT", "ADMIN_CORRECTION")]

    detail = client.get(f"{API}/attendance/sessions/{old['id']}", headers=admin).json()
    assert detail["status"] == "VOIDED"
    assert [e["source"] for e in detail["events"]] == ["KIOSK"]  # the original event is kept, never edited

    # a voided session cannot be corrected again, and the person can scan in again (no open session is left)
    assert correct(client, admin, old["id"]).status_code == 409
    assert scan(client, kiosk, person) == "ENTRY"

    audit = client.get(f"{API}/audit-logs", params={"entity_type": "attendance_session", "entity_id": old["id"],
                                                    "action": "CORRECT"}, headers=admin).json()
    assert len(audit) == 1
    assert audit[0]["reason"] == "Forgot to scan out"
    assert audit[0]["old_values"]["status"] == "OPEN" and audit[0]["new_values"]["status"] == "CLOSED"
    assert audit[0]["actor_id"] and audit[0]["ip_address"]


def test_session_list_filters(client, admin, kiosk, make_person):
    person = make_person()
    s = open_session(client, admin, kiosk, person)
    base = {"person_id": person["id"]}
    assert len(client.get(f"{API}/attendance/sessions", params={**base, "status": "OPEN"}, headers=admin).json()) == 1
    assert client.get(f"{API}/attendance/sessions", params={**base, "status": "CLOSED"}, headers=admin).json() == []
    assert client.get(f"{API}/attendance/sessions", params={**base, "type": "STUDENT"}, headers=admin).json() == []
    far = {"date_from": "2000-01-01", "date_to": "2000-12-31"}
    assert client.get(f"{API}/attendance/sessions", params={**base, **far}, headers=admin).json() == []
    assert client.get(f"{API}/attendance/sessions", params={"status": "BOGUS"}, headers=admin).status_code == 422
    assert client.get(f"{API}/attendance/sessions/{s['id']}", headers=admin).json()["events"][0]["event_type"] == "ENTRY"
    assert client.get(f"{API}/attendance/sessions/{uuid.uuid4()}", headers=admin).status_code == 404


def test_audit_log_trail(client, admin, make_person):
    p = make_person()
    client.patch(f"{API}/people/{p['id']}", json={"first_name": "Changed"}, headers=admin)
    client.post(f"{API}/people/{p['id']}/qr/rotate", headers=admin)
    rows = client.get(f"{API}/audit-logs", params={"entity_type": "person", "entity_id": p["id"]}, headers=admin).json()
    assert sorted(r["action"] for r in rows) == ["CREATE", "QR_ROTATE", "UPDATE"]
    update = next(r for r in rows if r["action"] == "UPDATE")
    assert update["old_values"]["first_name"] == "Test" and update["new_values"]["first_name"] == "Changed"
    assert "tt_" not in str(rows)  # raw QR tokens never reach the audit log

    assert client.get(f"{API}/audit-logs", params={"action": "BOGUS"}, headers=admin).status_code == 422
    assert client.get(f"{API}/audit-logs", params={"date_from": "2000-01-01", "date_to": "2000-01-02"},
                      headers=admin).json() == []


def test_reports_summary_and_exports(client, admin, kiosk, make_person, instant_type):
    # A person of their own type, so the report rows are exactly this person's.
    person = make_person(person_type=instant_type, first_name="=cmd|calc", last_name=uid("Rep"), group_name="Science")
    voided = open_session(client, admin, kiosk, person)
    assert correct(client, admin, voided["id"]).status_code == 201  # becomes a CLOSED 2 h session in March
    assert scan(client, kiosk, person) == "ENTRY"                   # a new session, still open today

    march = {"date_from": "2026-03-01", "date_to": "2026-03-31", "type": instant_type}
    rows = client.get(f"{API}/reports/summary", params=march, headers=admin).json()
    assert len(rows) == 1
    assert rows[0]["total_hours"] == 2.0 and rows[0]["days_present"] == 1
    assert rows[0]["sessions_closed"] == 1 and rows[0]["sessions_open"] == 0  # the voided session counts nowhere

    today = {"date_from": str(date.today()), "date_to": str(date.today()), "type": instant_type}
    rows_today = client.get(f"{API}/reports/summary", params=today, headers=admin).json()
    assert len(rows_today) == 1 and rows_today[0]["sessions_open"] == 1 and rows_today[0]["total_hours"] == 0

    bad_range = {"date_from": "2026-03-31", "date_to": "2026-03-01"}
    assert client.get(f"{API}/reports/summary", params=bad_range, headers=admin).status_code == 422

    csv_resp = client.get(f"{API}/reports/export", params={**march, "format": "csv"}, headers=admin)
    assert csv_resp.status_code == 200 and "attachment" in csv_resp.headers["content-disposition"]
    lines = csv_resp.content.decode("utf-8-sig").splitlines()
    assert lines[0] == ("Code,Name,Type,Group,Total hours,Days present,Late entries,"
                        "Closed sessions,Sessions still open")
    assert lines[1].split(",")[1].startswith("'=cmd|calc")  # formula injection neutralised

    xlsx_resp = client.get(f"{API}/reports/export", params={**march, "format": "xlsx"}, headers=admin)
    ws = load_workbook(io.BytesIO(xlsx_resp.content)).active
    assert ws["A1"].value == "Code" and ws["A2"].value == person["external_code"] and ws["E2"].value == 2.0
    assert ws["B2"].value.startswith("'=")

    assert client.get(f"{API}/reports/export", params={**march, "format": "pdf"}, headers=admin).status_code == 422

    exports = client.get(f"{API}/audit-logs", params={"action": "EXPORT"}, headers=admin).json()
    assert {"csv", "xlsx"} <= {e["new_values"]["format"] for e in exports}
